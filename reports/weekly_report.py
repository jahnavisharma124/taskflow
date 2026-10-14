"""
TaskFlow weekly report generator.

Reads project data straight from the TaskFlow MySQL database and writes an
Excel report (summary, hours per member, sprint progress, overdue tasks,
and every time entry in the period).

Examples (run from the taskflow/reports folder):
    python weekly_report.py --list
    python weekly_report.py --project 1
    python weekly_report.py --project 1 --days 14
"""
import argparse
import os
import sys
from datetime import date, timedelta
from pathlib import Path

import mysql.connector
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Reuse the backend's .env so the database password is stored in one place only.
ENV_PATH = Path(__file__).resolve().parent.parent / "backend" / ".env"
load_dotenv(ENV_PATH)

STATUS_LABEL = {"todo": "To Do", "in_progress": "In Progress", "review": "Review", "done": "Done"}


# ---------------------------------------------------------------- database ---
def connect():
    missing = [k for k in ("DB_HOST", "DB_USER", "DB_NAME") if not os.getenv(k)]
    if missing:
        sys.exit(f"Missing {', '.join(missing)}. Expected a backend/.env file at: {ENV_PATH}")
    try:
        return mysql.connector.connect(
            host=os.getenv("DB_HOST"),
            user=os.getenv("DB_USER"),
            password=os.getenv("DB_PASSWORD", ""),
            database=os.getenv("DB_NAME"),
        )
    except mysql.connector.Error as err:
        sys.exit(f"Could not connect to MySQL: {err}")


def query(conn, sql, params=()):
    cur = conn.cursor(dictionary=True)
    cur.execute(sql, params)
    rows = cur.fetchall()
    cur.close()
    return rows


def num(value):
    """MySQL returns DECIMAL as Decimal; convert to plain float for Excel."""
    return float(value) if value is not None else 0.0


def list_projects(conn):
    return query(conn, "SELECT id, name FROM projects ORDER BY id")


def collect_data(conn, project_id, start, end):
    projects = query(conn, "SELECT id, name FROM projects WHERE id = %s", (project_id,))
    if not projects:
        sys.exit(f"No project with id {project_id}. Run with --list to see project ids.")

    tasks = query(
        conn,
        """SELECT COUNT(*) AS total,
                  COALESCE(SUM(status = 'done'), 0) AS done,
                  COALESCE(SUM(status = 'in_progress'), 0) AS in_progress,
                  COALESCE(SUM(status = 'review'), 0) AS review,
                  COALESCE(SUM(status = 'todo'), 0) AS todo
           FROM tasks WHERE project_id = %s""",
        (project_id,),
    )[0]

    period_hours = query(
        conn,
        """SELECT COALESCE(SUM(tl.hours), 0) AS hours
           FROM time_logs tl JOIN tasks t ON t.id = tl.task_id
           WHERE t.project_id = %s AND tl.log_date BETWEEN %s AND %s""",
        (project_id, start, end),
    )[0]["hours"]

    members = query(
        conn,
        """SELECT u.name,
                  COALESCE(SUM(CASE WHEN tl.log_date BETWEEN %s AND %s THEN tl.hours END), 0) AS hours_period,
                  COALESCE(SUM(tl.hours), 0) AS hours_total
           FROM project_members pm
           JOIN users u ON u.id = pm.user_id
           LEFT JOIN tasks t ON t.project_id = pm.project_id
           LEFT JOIN time_logs tl ON tl.task_id = t.id AND tl.user_id = u.id
           WHERE pm.project_id = %s
           GROUP BY u.id, u.name
           ORDER BY hours_period DESC, u.name""",
        (start, end, project_id),
    )

    sprints = query(
        conn,
        """SELECT s.name, s.start_date, s.end_date,
                  COUNT(t.id) AS task_count,
                  COALESCE(SUM(t.status = 'done'), 0) AS done_count,
                  (SELECT COALESCE(SUM(tl.hours), 0)
                     FROM time_logs tl JOIN tasks t2 ON t2.id = tl.task_id
                    WHERE t2.sprint_id = s.id) AS hours
           FROM sprints s LEFT JOIN tasks t ON t.sprint_id = s.id
           WHERE s.project_id = %s
           GROUP BY s.id, s.name, s.start_date, s.end_date
           ORDER BY s.start_date""",
        (project_id,),
    )

    overdue = query(
        conn,
        """SELECT t.title, t.priority, t.status, t.due_date, u.name AS assignee,
                  DATEDIFF(CURDATE(), t.due_date) AS days_overdue
           FROM tasks t LEFT JOIN users u ON u.id = t.assignee_id
           WHERE t.project_id = %s AND t.due_date < CURDATE() AND t.status <> 'done'
           ORDER BY t.due_date""",
        (project_id,),
    )

    entries = query(
        conn,
        """SELECT tl.log_date, u.name AS member, t.title AS task, tl.hours, tl.note
           FROM time_logs tl
           JOIN tasks t ON t.id = tl.task_id
           JOIN users u ON u.id = tl.user_id
           WHERE t.project_id = %s AND tl.log_date BETWEEN %s AND %s
           ORDER BY tl.log_date, u.name""",
        (project_id, start, end),
    )

    return {
        "project": projects[0],
        "start": start,
        "end": end,
        "tasks": {k: int(v) for k, v in tasks.items()},
        "period_hours": num(period_hours),
        "members": members,
        "sprints": sprints,
        "overdue": overdue,
        "entries": entries,
    }


# ------------------------------------------------------------------- excel ---
HEADER_FILL = PatternFill("solid", fgColor="2563EB")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_table(ws, start_row, headers, rows):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left")
    for r, row in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    return start_row + len(rows) + 1


def autosize(ws):
    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 2, 10), 60)


def pct(done, total):
    return round(100 * done / total, 1) if total else 0.0


def build_workbook(data, path):
    wb = Workbook()

    # Summary
    ws = wb.active
    ws.title = "Summary"
    t = data["tasks"]
    ws["A1"] = f"TaskFlow report: {data['project']['name']}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Period: {data['start']} to {data['end']}"
    ws["A3"] = f"Generated: {date.today()}"
    write_table(
        ws,
        5,
        ["Metric", "Value"],
        [
            ["Total tasks", t["total"]],
            ["To Do", t["todo"]],
            ["In Progress", t["in_progress"]],
            ["Review", t["review"]],
            ["Done", t["done"]],
            ["Completion %", pct(t["done"], t["total"])],
            ["Overdue tasks", len(data["overdue"])],
            ["Hours logged in period", round(data["period_hours"], 2)],
        ],
    )
    autosize(ws)

    # Hours by member
    ws = wb.create_sheet("Hours by Member")
    write_table(
        ws,
        1,
        ["Member", "Hours (period)", "Hours (all time)"],
        [[m["name"], round(num(m["hours_period"]), 2), round(num(m["hours_total"]), 2)] for m in data["members"]],
    )
    autosize(ws)

    # Sprints
    ws = wb.create_sheet("Sprints")
    write_table(
        ws,
        1,
        ["Sprint", "Start", "End", "Tasks", "Done", "Completion %", "Hours logged"],
        [
            [
                s["name"],
                str(s["start_date"] or ""),
                str(s["end_date"] or ""),
                int(s["task_count"]),
                int(s["done_count"]),
                pct(int(s["done_count"]), int(s["task_count"])),
                round(num(s["hours"]), 2),
            ]
            for s in data["sprints"]
        ],
    )
    autosize(ws)

    # Overdue
    ws = wb.create_sheet("Overdue Tasks")
    write_table(
        ws,
        1,
        ["Task", "Assignee", "Priority", "Status", "Due date", "Days overdue"],
        [
            [
                o["title"],
                o["assignee"] or "Unassigned",
                o["priority"],
                STATUS_LABEL.get(o["status"], o["status"]),
                str(o["due_date"]),
                int(o["days_overdue"]),
            ]
            for o in data["overdue"]
        ],
    )
    autosize(ws)

    # Time entries
    ws = wb.create_sheet("Time Entries")
    write_table(
        ws,
        1,
        ["Date", "Member", "Task", "Hours", "Note"],
        [[str(e["log_date"]), e["member"], e["task"], round(num(e["hours"]), 2), e["note"] or ""] for e in data["entries"]],
    )
    autosize(ws)

    wb.save(path)


# -------------------------------------------------------------------- main ---
def main():
    parser = argparse.ArgumentParser(description="Generate a TaskFlow Excel report.")
    parser.add_argument("--list", action="store_true", help="list projects and their ids, then exit")
    parser.add_argument("--project", type=int, help="project id to report on")
    parser.add_argument("--days", type=int, default=7, help="how many days back to cover (default 7)")
    parser.add_argument("--out", help="output folder (default: reports/output)")
    args = parser.parse_args()

    conn = connect()
    try:
        if args.list:
            for p in list_projects(conn):
                print(f"{p['id']}: {p['name']}")
            return
        if args.project is None:
            parser.error("--project is required (use --list to see ids)")
        if args.days < 1:
            parser.error("--days must be at least 1")

        end = date.today()
        start = end - timedelta(days=args.days - 1)  # --days 7 means today and the 6 days before
        data = collect_data(conn, args.project, start, end)
    finally:
        conn.close()

    out_dir = Path(args.out) if args.out else Path(__file__).resolve().parent / "output"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"taskflow_report_project{args.project}_{end}.xlsx"
    build_workbook(data, path)
    print(f"Report saved: {path}")


if __name__ == "__main__":
    main()
